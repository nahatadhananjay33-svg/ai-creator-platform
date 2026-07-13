"""Media database: SQLite source of truth + CSV/XLSX exports."""
from __future__ import annotations

import csv
import sqlite3
from pathlib import Path
from typing import List, Optional

from .models import MEDIA_COLUMNS, MediaRecord

_SQL_TYPES = {"id": "INTEGER PRIMARY KEY AUTOINCREMENT", "duration": "REAL", "fps": "REAL"}
_INSERT_COLS = [c for c in MEDIA_COLUMNS if c != "id"]


class MediaDB:
    """Upsert-by-(platform, item_id) store; ids assigned by SQLite."""

    def __init__(self, path: Path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._init()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(str(self.path))
        conn.row_factory = sqlite3.Row
        return conn

    def _init(self) -> None:
        cols = ", ".join(f"{c} {_SQL_TYPES.get(c, 'TEXT')}" for c in MEDIA_COLUMNS)
        conn = self._connect()
        try:
            conn.execute(f"CREATE TABLE IF NOT EXISTS media ({cols})")
            conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_item "
                         "ON media(platform, item_id)")
            conn.commit()
        finally:
            conn.close()

    def get(self, platform: str, item_id: str) -> Optional[dict]:
        conn = self._connect()
        try:
            row = conn.execute("SELECT * FROM media WHERE platform=? AND item_id=?",
                               (platform, item_id)).fetchone()
            return dict(row) if row else None
        finally:
            conn.close()

    def get_by_checksum(self, checksum: str) -> Optional[dict]:
        if not checksum:
            return None
        conn = self._connect()
        try:
            row = conn.execute("SELECT * FROM media WHERE checksum=?", (checksum,)).fetchone()
            return dict(row) if row else None
        finally:
            conn.close()

    def upsert(self, rec: MediaRecord) -> None:
        d = rec.as_dict()
        placeholders = ", ".join("?" for _ in _INSERT_COLS)
        updates = ", ".join(f"{c}=excluded.{c}" for c in _INSERT_COLS)
        conn = self._connect()
        try:
            conn.execute(
                f"INSERT INTO media ({', '.join(_INSERT_COLS)}) VALUES ({placeholders}) "
                f"ON CONFLICT(platform, item_id) DO UPDATE SET {updates}",
                [d[c] for c in _INSERT_COLS],
            )
            conn.commit()
        finally:
            conn.close()

    def all(self) -> List[dict]:
        conn = self._connect()
        try:
            return [dict(r) for r in conn.execute(
                "SELECT * FROM media ORDER BY platform, item_id")]
        finally:
            conn.close()

    def downloaded_hours(self, platform: str) -> float:
        conn = self._connect()
        try:
            v = conn.execute(
                "SELECT COALESCE(SUM(duration), 0) FROM media "
                "WHERE platform=? AND status='downloaded'", (platform,)).fetchone()[0]
            return float(v or 0.0) / 3600.0
        finally:
            conn.close()

    # --- exports --------------------------------------------------------------
    def export_csv(self, path: Path) -> Path:
        path = Path(path)
        rows = self.all()
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(MEDIA_COLUMNS)
            for r in rows:
                w.writerow([r.get(c, "") for c in MEDIA_COLUMNS])
        return path

    def export_xlsx(self, path: Path) -> Path:
        from openpyxl import Workbook
        path = Path(path)
        wb = Workbook()
        ws = wb.active
        ws.title = "media"
        ws.append(MEDIA_COLUMNS)
        for r in self.all():
            ws.append([r.get(c, "") for c in MEDIA_COLUMNS])
        wb.save(str(path))
        return path
