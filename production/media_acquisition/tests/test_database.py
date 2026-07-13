"""Media database upsert / query / export."""
from __future__ import annotations

import csv
import sqlite3

from production.media_acquisition.database import MediaDB
from production.media_acquisition.models import MEDIA_COLUMNS, MediaRecord


def _rec(item_id="a", status="downloaded", checksum="abc", duration=3600.0):
    return MediaRecord(
        id=0, platform="youtube", url=f"https://x/{item_id}", title="t",
        publish_date="2026-01-01", duration=duration, resolution="1080p", fps=30.0,
        checksum=checksum, filename=f"{item_id}.mp4", status=status,
        download_time="2026-01-01T00:00:00+00:00", item_id=item_id, codec="h264",
        kind="long_form")


def test_upsert_is_idempotent(tmp_path):
    db = MediaDB(tmp_path / "m.sqlite")
    db.upsert(_rec("a"))
    db.upsert(_rec("a", status="downloaded"))   # same key -> update, no dup row
    rows = db.all()
    assert len(rows) == 1
    assert rows[0]["item_id"] == "a" and rows[0]["id"] == 1


def test_get_and_checksum_and_hours(tmp_path):
    db = MediaDB(tmp_path / "m.sqlite")
    db.upsert(_rec("a", checksum="c1", duration=1800.0))
    db.upsert(_rec("b", checksum="c2", duration=1800.0))
    assert db.get("youtube", "a")["checksum"] == "c1"
    assert db.get_by_checksum("c2")["item_id"] == "b"
    assert abs(db.downloaded_hours("youtube") - 1.0) < 1e-6   # 2x1800s = 1h


def test_exports(tmp_path):
    db = MediaDB(tmp_path / "m.sqlite")
    db.upsert(_rec("a"))
    db.upsert(_rec("b"))
    csv_path = db.export_csv(tmp_path / "media.csv")
    xlsx_path = db.export_xlsx(tmp_path / "media.xlsx")
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows[0] == MEDIA_COLUMNS and len(rows) == 3
    from openpyxl import load_workbook
    ws = load_workbook(xlsx_path).active
    assert [c.value for c in ws[1]] == MEDIA_COLUMNS and ws.max_row == 3
