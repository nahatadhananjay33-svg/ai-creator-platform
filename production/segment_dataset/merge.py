"""Phase V4.1 STEPS 6-7 + validation: SHA256-deduped merge and comparison.

Only ACCEPTED segments from the new (temporary) dataset are merged into the
production dataset; duplicates are detected by content hash against every
accepted WAV already in production. The production dataset is otherwise
untouched — validation proves it (row identity for the old prefix, file
existence, sqlite/csv/xlsx row agreement, hash uniqueness).
"""
from __future__ import annotations

import csv
import shutil
from collections import Counter
from pathlib import Path
from typing import Dict, List, Tuple

from production.media_acquisition.download import sha256

from .models import SegmentRecord
from .storage import load_segments, row_to_record, write_all


def _accepted_hashes(rows: List[dict]) -> Dict[str, str]:
    """sha256 -> segment_file for every accepted row whose WAV exists."""
    out: Dict[str, str] = {}
    for r in rows:
        if r["accepted"] and Path(r["audio_path"]).exists():
            out[sha256(Path(r["audio_path"]))] = r["segment_file"]
    return out


def merge_accepted(new_dir: Path, prod_dir: Path) -> dict:
    """Merge new accepted segments into production; returns merge stats."""
    new_dir, prod_dir = Path(new_dir), Path(prod_dir)
    prod_meta = prod_dir / "metadata"
    prod_acc = prod_dir / "accepted_segments"
    prod_acc.mkdir(parents=True, exist_ok=True)

    old_rows = load_segments(prod_meta / "dataset.sqlite")
    new_rows = load_segments(new_dir / "metadata" / "dataset.sqlite")
    existing = _accepted_hashes(old_rows)

    added: List[SegmentRecord] = []
    duplicates: List[Tuple[str, str]] = []       # (new file, duplicate-of)
    missing: List[str] = []
    for r in sorted((r for r in new_rows if r["accepted"]),
                    key=lambda r: r["segment_file"]):
        src = Path(r["audio_path"])
        if not src.exists():
            missing.append(r["segment_file"])
            continue
        digest = sha256(src)
        if digest in existing:
            duplicates.append((r["segment_file"], existing[digest]))
            continue
        target = prod_acc / r["segment_file"]
        if target.exists():                      # same name, different content
            target = prod_acc / (target.stem + "__v41" + target.suffix)
        shutil.copyfile(src, target)
        rec = row_to_record(r)
        rec.segment_file = target.name
        rec.audio_path = str(target)
        added.append(rec)
        existing[digest] = target.name

    merged = [row_to_record(r) for r in old_rows] + added
    for i, rec in enumerate(merged, 1):
        rec.id = i
    write_all(merged, prod_meta)
    return {"old_rows": len(old_rows), "new_accepted": len(added) + len(duplicates),
            "added": len(added), "duplicates_skipped": len(duplicates),
            "missing_files": missing, "merged_rows": len(merged)}


# --- STEP 7: comparison -------------------------------------------------------------

def _metrics(rows: List[dict]) -> dict:
    acc = [r for r in rows if r["accepted"]]
    n = len(acc)
    return {
        "accepted_clips": n,
        "accepted_hours": round(sum(r["duration"] for r in acc) / 3600.0, 3),
        "avg_snr_db": round(sum(r["snr_db"] for r in acc) / n, 2) if n else 0.0,
        "avg_loudness_dbfs": round(sum(r["loudness_dbfs"] for r in acc) / n, 2) if n else 0.0,
        "avg_duration_s": round(sum(r["duration"] for r in acc) / n, 1) if n else 0.0,
        "quality_distribution": dict(Counter(r["quality"] for r in acc)),
        "rejection_reasons": Counter(r["reason"] for r in rows
                                     if not r["accepted"]).most_common(5),
        "recovery_pct": round(100.0 * n / len(rows), 1) if rows else 0.0,
    }


def comparison(old_rows: List[dict], new_rows: List[dict],
               merged_rows: List[dict]) -> Tuple[dict, str]:
    m = {"old": _metrics(old_rows), "new": _metrics(new_rows),
         "merged": _metrics(merged_rows)}
    net_h = m["merged"]["accepted_hours"] - m["old"]["accepted_hours"]
    rows = [("Accepted clips", "accepted_clips", "{:.0f}"),
            ("Accepted hours", "accepted_hours", "{:.3f}"),
            ("Average SNR (dB)", "avg_snr_db", "{:.2f}"),
            ("Average loudness (dBFS)", "avg_loudness_dbfs", "{:.2f}"),
            ("Average duration (s)", "avg_duration_s", "{:.1f}"),
            ("Recovery (accepted %)", "recovery_pct", "{:.1f}")]
    lines = [f"  {'metric':<26} {'Old (V4)':>12} {'New (V4.1)':>12} {'Merged':>12}",
             "  " + "-" * 66]
    for label, key, fmt in rows:
        lines.append(f"  {label:<26} {fmt.format(m['old'][key]):>12} "
                     f"{fmt.format(m['new'][key]):>12} {fmt.format(m['merged'][key]):>12}")
    lines.append("")
    for name in ("old", "new", "merged"):
        q = ", ".join(f"{k}: {v}" for k, v in sorted(m[name]["quality_distribution"].items()))
        lines.append(f"  Quality ({name:<6}) : {q or 'n/a'}")
    lines.append("")
    lines.append("  Top rejection reasons (new run):")
    for reason, count in m["new"]["rejection_reasons"] or [("none", 0)]:
        lines.append(f"    - {reason} ({count})")
    lines.append("")
    lines.append(f"  Net increase: +{net_h:.3f} h accepted "
                 f"(+{m['merged']['accepted_clips'] - m['old']['accepted_clips']} clips)")
    m["net_increase_hours"] = round(net_h, 3)
    return m, "\n".join(lines)


# --- validation -----------------------------------------------------------------------

def _csv_rows(path: Path) -> int:
    with open(path, newline="", encoding="utf-8") as f:
        return sum(1 for _ in csv.reader(f)) - 1               # minus header


def _xlsx_rows(path: Path) -> int:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True)
    n = wb.active.max_row - 1
    wb.close()
    return n


def validate_merge(prod_dir: Path, old_rows: List[dict], added: int,
                   audit_ok: bool) -> Dict[str, bool]:
    """The spec's validation checklist, each item measured."""
    prod_dir = Path(prod_dir)
    meta = prod_dir / "metadata"
    rows = load_segments(meta / "dataset.sqlite")

    old_keys = [(r["source_kind"], r["segment_file"]) for r in old_rows]
    new_keys = [(r["source_kind"], r["segment_file"]) for r in rows[:len(old_rows)]]
    accepted = [r for r in rows if r["accepted"]]
    hashes = [sha256(Path(r["audio_path"])) for r in accepted
              if Path(r["audio_path"]).exists()]

    return {
        "existing dataset unchanged except additions":
            new_keys == old_keys and len(rows) == len(old_rows) + added,
        "no duplicate segments (sha256)": len(hashes) == len(set(hashes)),
        "metadata synchronized (ids sequential)":
            [r["id"] for r in rows] == list(range(1, len(rows) + 1)),
        "sqlite consistent": len(rows) == len(old_rows) + added,
        "csv consistent": _csv_rows(meta / "dataset.csv") == len(rows),
        "xlsx consistent": _xlsx_rows(meta / "dataset.xlsx") == len(rows),
        "all accepted files exist":
            all(Path(r["audio_path"]).exists() for r in accepted),
        "random audit passed": audit_ok,
    }


def format_validation(checks: Dict[str, bool]) -> str:
    return "\n".join(f"  {'[OK]  ' if ok else '[FAIL]'} {name}"
                     for name, ok in checks.items())
