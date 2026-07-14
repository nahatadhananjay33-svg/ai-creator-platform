"""Persist segment records to SQLite, CSV, and XLSX.

Mirrors ``production.voice_dataset.storage`` (deterministic full rewrite per
run) for the segment schema.
"""
from __future__ import annotations

import csv
import sqlite3
from pathlib import Path
from typing import List

from .models import SEGMENT_COLUMNS, SegmentRecord

_SQL_TYPES = {
    "id": "INTEGER PRIMARY KEY", "start_s": "REAL", "end_s": "REAL",
    "duration": "REAL", "speech_duration": "REAL", "snr_db": "REAL",
    "loudness_dbfs": "REAL", "noise_floor_dbfs": "REAL",
    "multi_speaker_confidence": "REAL", "f0_iqr_hz": "REAL",
    "accepted": "INTEGER", "has_music": "INTEGER", "sample_rate": "INTEGER",
}
_BOOL_COLS = {"accepted", "has_music"}


def _row(rec: SegmentRecord) -> list:
    d = rec.as_dict()
    return [int(d[c]) if c in _BOOL_COLS else d[c] for c in SEGMENT_COLUMNS]


def load_segments(path: Path) -> List[dict]:
    path = Path(path)
    if not path.exists():
        return []
    conn = sqlite3.connect(str(path))
    conn.row_factory = sqlite3.Row
    try:
        rows = [dict(r) for r in conn.execute("SELECT * FROM dataset")]
    finally:
        conn.close()
    for r in rows:
        r["accepted"] = bool(r.get("accepted"))
        r["has_music"] = bool(r.get("has_music"))
    return rows


def row_to_record(row: dict) -> SegmentRecord:
    return SegmentRecord(**{c: row[c] for c in SEGMENT_COLUMNS})


def write_sqlite(records: List[SegmentRecord], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        path.unlink()
    cols = ", ".join(f"{c} {_SQL_TYPES.get(c, 'TEXT')}" for c in SEGMENT_COLUMNS)
    placeholders = ", ".join("?" for _ in SEGMENT_COLUMNS)
    conn = sqlite3.connect(str(path))
    try:
        conn.execute(f"CREATE TABLE dataset ({cols})")
        conn.executemany(
            f"INSERT INTO dataset ({', '.join(SEGMENT_COLUMNS)}) VALUES ({placeholders})",
            [_row(r) for r in records])
        conn.commit()
    finally:
        conn.close()
    return path


def write_csv(records: List[SegmentRecord], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(SEGMENT_COLUMNS)
        for r in records:
            w.writerow(_row(r))
    return path


def write_xlsx(records: List[SegmentRecord], path: Path) -> Path:
    from openpyxl import Workbook
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "dataset"
    ws.append(SEGMENT_COLUMNS)
    for r in records:
        ws.append(_row(r))
    wb.save(str(path))
    return path


def write_all(records: List[SegmentRecord], metadata_dir: Path) -> None:
    metadata_dir = Path(metadata_dir)
    write_sqlite(records, metadata_dir / "dataset.sqlite")
    write_csv(records, metadata_dir / "dataset.csv")
    write_xlsx(records, metadata_dir / "dataset.xlsx")
