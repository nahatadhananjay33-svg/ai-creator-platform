"""Persist dataset records to SQLite, CSV, and XLSX (deterministic)."""
from __future__ import annotations

import csv
import sqlite3
from pathlib import Path
from typing import List

from .models import RECORD_COLUMNS, Record

# SQLite column types keyed by name (booleans stored as 0/1 INTEGER).
_SQL_TYPES = {
    "id": "INTEGER PRIMARY KEY", "duration": "REAL", "speech_duration": "REAL",
    "sample_rate": "INTEGER", "channels": "INTEGER", "bitrate": "INTEGER",
    "silence_pct": "REAL", "loudness_dbfs": "REAL", "snr_db": "REAL",
    "accepted": "INTEGER", "has_speech": "INTEGER", "multi_speaker": "INTEGER",
    "has_music": "INTEGER",
}
_BOOL_COLS = {"accepted", "has_speech", "multi_speaker", "has_music"}


def _row(rec: Record) -> list:
    d = rec.as_dict()
    return [int(d[c]) if c in _BOOL_COLS else d[c] for c in RECORD_COLUMNS]


def write_sqlite(records: List[Record], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        path.unlink()                      # deterministic: rebuild each run
    cols = ", ".join(f"{c} {_SQL_TYPES.get(c, 'TEXT')}" for c in RECORD_COLUMNS)
    placeholders = ", ".join("?" for _ in RECORD_COLUMNS)
    conn = sqlite3.connect(str(path))
    try:
        conn.execute(f"CREATE TABLE dataset ({cols})")
        conn.executemany(
            f"INSERT INTO dataset ({', '.join(RECORD_COLUMNS)}) VALUES ({placeholders})",
            [_row(r) for r in records],
        )
        conn.commit()
    finally:
        conn.close()
    return path


def write_csv(records: List[Record], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(RECORD_COLUMNS)
        for r in records:
            w.writerow(_row(r))
    return path


def write_xlsx(records: List[Record], path: Path) -> Path:
    try:
        from openpyxl import Workbook
    except ImportError as e:                # pragma: no cover
        raise RuntimeError("openpyxl is required for XLSX output: pip install openpyxl") from e
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "dataset"
    ws.append(RECORD_COLUMNS)
    for r in records:
        ws.append(_row(r))
    wb.save(str(path))
    return path
