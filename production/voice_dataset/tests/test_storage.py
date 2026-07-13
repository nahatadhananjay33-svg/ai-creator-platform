"""SQLite / CSV / XLSX writers round-trip the record columns."""
from __future__ import annotations

import csv
import sqlite3

from production.voice_dataset.models import RECORD_COLUMNS, Record
from production.voice_dataset.storage import write_csv, write_sqlite, write_xlsx


def _records():
    return [
        Record(id=1, platform="youtube", source="youtube/a.wav", filename="a.wav",
               duration=12.0, speech_duration=9.0, classification="talking_head",
               quality="good", accepted=True, reason="ok", audio_path="/x/a.wav",
               sample_rate=16000, channels=1, bitrate=256000, silence_pct=0.25,
               loudness_dbfs=-20.0, snr_db=22.0, has_speech=True, multi_speaker=False,
               has_music=False, noise_estimate="low"),
        Record(id=2, platform="instagram", source="instagram/b.wav", filename="b.wav",
               duration=5.0, speech_duration=0.0, classification="unknown",
               quality="poor", accepted=False, reason="no speech detected",
               audio_path="/x/b.wav"),
    ]


def test_sqlite(tmp_path):
    recs = _records()
    db = write_sqlite(recs, tmp_path / "d.sqlite")
    conn = sqlite3.connect(str(db))
    cols = [r[1] for r in conn.execute("PRAGMA table_info(dataset)")]
    assert cols == RECORD_COLUMNS
    rows = conn.execute("SELECT id, accepted FROM dataset ORDER BY id").fetchall()
    assert rows == [(1, 1), (2, 0)]        # booleans stored as int
    conn.close()


def test_csv(tmp_path):
    recs = _records()
    path = write_csv(recs, tmp_path / "d.csv")
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows[0] == RECORD_COLUMNS
    assert len(rows) == 3                   # header + 2
    assert rows[1][rows[0].index("filename")] == "a.wav"


def test_xlsx(tmp_path):
    from openpyxl import load_workbook
    recs = _records()
    path = write_xlsx(recs, tmp_path / "d.xlsx")
    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    assert header == RECORD_COLUMNS
    assert ws.max_row == 3
